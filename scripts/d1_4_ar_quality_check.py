"""1일차 d1-4: 영업팀 미수금 엑셀(d1-4_ar_sales_team_2026-09.xlsx) 품질 점검.

원본은 읽기만 한다(openpyxl load_workbook, 저장하지 않음).
실행: uv run python scripts/d1_4_ar_quality_check.py
결과는 화면에 출력하고, 보고서(workbench/day1/outputs/d1-4_quality_report.md)의 숫자 근거로 쓴다.
"""
import re
import statistics
import sys
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path

import openpyxl

sys.stdout.reconfigure(encoding="utf-8")

SRC = Path("data/day1/d1-4_ar_sales_team_2026-09.xlsx")
AS_OF = date(2026, 9, 30)  # 요약 시트 A10의 기준일
DETAIL_SHEETS = ["아시아", "유럽·미주", "중동·아프리카"]
COL = dict(owner=1, buyer=2, country=3, terms=4, inv=5, ship=6, due=7, krw=8, fx=9, ccy=10, memo=11)


def parse_date(v, inv_no):
    """날짜 셀 → (date, 형식). 연도 없는 'M월 D일'은 인보이스 번호의 연도(INV-YYMM)로 보충."""
    if isinstance(v, datetime):
        return v.date(), "날짜형"
    if isinstance(v, str):
        m = re.fullmatch(r"(\d{2})\.(\d{1,2})\.(\d{1,2})", v)
        if m:
            return date(2000 + int(m[1]), int(m[2]), int(m[3])), "문자 YY.M.D"
        m = re.fullmatch(r"(\d{1,2})월 (\d{1,2})일", v)
        if m:
            yy = int(inv_no[4:6])  # INV-2608-... → 26
            return date(2000 + yy, int(m[1]), int(m[2])), "문자 M월 D일(연도 없음)"
    return None, f"알 수 없음: {v!r}"


def term_days(terms):
    """결제조건 문자열에서 만기 일수 추출(없으면 None)."""
    if re.search(r"at sight|일람불", terms, re.I):
        return "sight"
    if re.search(r"30/70|30%", terms):
        return 30
    m = re.search(r"(\d{2,3})", terms)
    return int(m[1]) if m else None


def term_family(terms):
    t = terms.upper().replace(" ", "")
    if "L/C" in t or t.startswith("LC"):
        return "L/C"
    if "D/P" in t:
        return "D/P"
    if "D/A" in t or t.startswith("DA"):
        return "D/A"
    if "30/70" in t or "30%" in t:
        return "T/T 30/70"
    if "O/A" in t or t.startswith("OA") or "사후송금" in terms:
        return "O/A"
    return "기타"


wb = openpyxl.load_workbook(SRC, data_only=False)  # 읽기 전용 용도. save() 호출하지 않음.

rows = []          # 인보이스 행
structure = {}     # 시트별 구조
for name in DETAIL_SHEETS:
    ws = wb[name]
    info = dict(merged=[str(r) for r in ws.merged_cells.ranges], subtotals=[], blanks=[], total=None,
                max_row=ws.max_row)
    for r in range(5, ws.max_row + 1):
        a = ws.cell(r, 1).value
        if a is None and all(ws.cell(r, c).value is None for c in range(1, 12)):
            info["blanks"].append(r)
        elif isinstance(a, str) and a.endswith("소계"):
            info["subtotals"].append((r, a, ws.cell(r, 8).value))
        elif a == "합계":
            info["total"] = (r, ws.cell(r, 8).value)
        else:
            rec = {k: ws.cell(r, c).value for k, c in COL.items()}
            rec.update(sheet=name, row=r)
            rows.append(rec)
    structure[name] = info

print("=" * 70, "\n[1] 시트 구조")
print("시트 목록:", wb.sheetnames)
for name, info in structure.items():
    n = sum(1 for x in rows if x["sheet"] == name)
    print(f"- {name}: 마지막 행 {info['max_row']}, 인보이스 {n}행, 소계 {len(info['subtotals'])}줄, "
          f"빈 줄 {info['blanks']}, 합계 {info['total']}")
    for s in info["subtotals"]:
        print("    소계", s)
ws = wb["요약"]
for r in range(1, ws.max_row + 1):
    vals = [ws.cell(r, c).value for c in range(1, 5)]
    if any(v is not None for v in vals):
        print(f"  요약 {r}행: {vals}")

print("=" * 70, "\n[2] 금액(원화) 문자열 — SUM에서 빠짐")
text_amt = [x for x in rows if isinstance(x["krw"], str)]
for x in text_amt:
    print(f"  {x['sheet']} H{x['row']} = {x['krw']!r}")
by_sheet = defaultdict(lambda: [0, 0])
for x in rows:
    v = x["krw"]
    num = int(v.replace(",", "")) if isinstance(v, str) else v
    by_sheet[x["sheet"]][0] += num if not isinstance(v, str) else 0   # 엑셀 SUM이 보는 값
    by_sheet[x["sheet"]][1] += num                                     # 문자를 숫자로 바꾼 값
for s, (excel_sum, true_sum) in by_sheet.items():
    print(f"  {s}: 엑셀 SUM 기준 {excel_sum:,} / 문자 포함 {true_sum:,} / 누락 {true_sum - excel_sum:,} 천원")
print(f"  전체: 엑셀 {sum(v[0] for v in by_sheet.values()):,} / 문자 포함 {sum(v[1] for v in by_sheet.values()):,}")

print("=" * 70, "\n[3] 환산율(원화천원×1000 ÷ 외화) 이상치")
CCY_NORM = {"usd": "USD", "US$": "USD"}
rates = defaultdict(list)
for x in rows:
    krw = int(x["krw"].replace(",", "")) if isinstance(x["krw"], str) else x["krw"]
    ccy = CCY_NORM.get(x["ccy"], x["ccy"])
    x["rate"] = krw * 1000 / x["fx"]
    x["ccy_n"] = ccy
    rates[ccy].append(x["rate"])
for ccy, rs in rates.items():
    med = statistics.median(rs)
    print(f"  {ccy}: {len(rs)}건, 중앙값 {med:,.2f}, 범위 {min(rs):,.2f} ~ {max(rs):,.2f}")
    for x in rows:
        if x["ccy_n"] == ccy and not (0.8 * med <= x["rate"] <= 1.2 * med):
            print(f"    이상: {x['sheet']} {x['row']}행 H={x['krw']} I={x['fx']} {x['ccy']} → {x['rate']:,.2f}"
                  f" (중앙값의 {x['rate'] / med:,.0f}배)")

print("=" * 70, "\n[4] 통화 · 국가 표기")
print("  통화:", Counter(x["ccy"] for x in rows))
print("  국가:", Counter(x["country"] for x in rows))
for v in ["usd", "US$"]:
    print(f"    통화 {v!r}:", [(x["sheet"], x["row"]) for x in rows if x["ccy"] == v])
for v in ["Viet Nam", "USA", "UK", "Turkey"]:
    print(f"    국가 {v!r}:", [(x["sheet"], x["row"]) for x in rows if x["country"] == v])

print("=" * 70, "\n[5] 결제조건 표기 변형")
fam = defaultdict(Counter)
for x in rows:
    fam[term_family(x["terms"])][x["terms"]] += 1
for f, c in fam.items():
    print(f"  {f}: {len(c)}가지 → {dict(c)}")
print("  같은 바이어에 조건 계열이 둘 이상:")
buyer_fam = defaultdict(set)
for x in rows:
    buyer_fam[x["buyer"]].add(term_family(x["terms"]))
for b, fs in buyer_fam.items():
    if len(fs) > 1:
        print(f"    {b}: {fs} →", [(x['sheet'], x['row'], x['terms']) for x in rows if x['buyer'] == b
                                   and term_family(x['terms']) != 'O/A'])

print("=" * 70, "\n[6] 날짜 형식")
fmt_cnt = Counter()
for x in rows:
    x["ship_d"], f1 = parse_date(x["ship"], x["inv"])
    x["due_d"], f2 = parse_date(x["due"], x["inv"])
    fmt_cnt[("선적일", f1)] += 1
    fmt_cnt[("만기일", f2)] += 1
for k, v in sorted(fmt_cnt.items()):
    print(f"  {k}: {v}")
for s in DETAIL_SHEETS:
    sub = [x for x in rows if x["sheet"] == s]
    c = Counter(parse_date(x["ship"], x["inv"])[1] for x in sub) + Counter(parse_date(x["due"], x["inv"])[1] for x in sub)
    print(f"  {s}: {dict(c)}")

print("=" * 70, "\n[7] 만기일 − 선적일 vs 결제조건 일수")
for x in rows:
    td = term_days(x["terms"])
    gap = (x["due_d"] - x["ship_d"]).days
    if td == "sight":
        print(f"  (일람불) {x['sheet']} {x['row']}행 {x['terms']}: 선적→만기 {gap}일")
    elif td is not None and gap != td:
        print(f"  불일치 {x['sheet']} {x['row']}행 {x['terms']}: {gap}일")
inv_month_mismatch = [(x["sheet"], x["row"], x["inv"], x["ship_d"]) for x in rows
                      if x["inv"][4:8] != x["ship_d"].strftime("%y%m")]
print("  인보이스번호 YYMM ≠ 선적 연월:", inv_month_mismatch)

print("=" * 70, "\n[8] 인보이스 번호 중복")
dup = Counter(x["inv"] for x in rows)
for inv, n in dup.items():
    if n > 1:
        for x in rows:
            if x["inv"] == inv:
                print(f"  {inv}: {x['sheet']} {x['row']}행 {x['owner']} {x['buyer']} {x['country']} H={x['krw']} I={x['fx']}")

print("=" * 70, "\n[9] 메모의 부도·법무 건, 만기 경과")
bad = [x for x in rows if x["memo"] and "부도" in x["memo"]]
for x in bad:
    print(f"  {x['sheet']} {x['row']}행 {x['buyer']} {x['inv']} 만기 {x['due_d']} H={x['krw']} | {x['memo']}")
bad_sum = sum(int(x["krw"].replace(",", "")) if isinstance(x["krw"], str) else x["krw"] for x in bad)
print(f"  부도 메모 {len(bad)}건, 원화 {bad_sum:,} 천원")
overdue = [x for x in rows if x["due_d"] <= AS_OF]
print(f"  기준일 {AS_OF} 이전·당일 만기(만기 경과) {len(overdue)}건 / 전체 {len(rows)}건")
print("  메모 있는 행:", sum(1 for x in rows if x["memo"]), "건")
print("  메모 값 종류:", Counter(x["memo"] for x in rows if x["memo"]))

print("=" * 70, "\n[10] 시트 지역 ↔ 국가 불일치")
REGION = {"아시아": {"Japan", "Indonesia", "Taiwan", "India", "China", "Thailand", "Vietnam", "Viet Nam"}}
for x in rows:
    if x["sheet"] != "아시아" and x["country"] in REGION["아시아"]:
        print(f"  {x['sheet']} {x['row']}행 {x['buyer']} {x['country']}")
